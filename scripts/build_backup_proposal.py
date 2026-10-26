#!/usr/bin/env python3
"""Build the executive backup-proposal document for บริษัท ยุธาภัคร์ จำกัด."""

from __future__ import annotations

import os
import re
import uuid
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Mm, Pt, RGBColor, Twips
from docx.table import _Cell
from fontTools.ttLib import TTFont

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "โครงการปรับปรุงระบบการสำรองข้อมูล.docx")

NAVY = "0E2A47"
NAVY_MID = "1A4066"
INK = "1B242E"
MUTED = "5C6770"
BRASS = "A68546"
BRASS_TEXT = "6E5324"
BRASS_ON_NAVY = "E7D7B0"
CREAM = "F7F4EE"
CREAM_DEEP = "EFE8DA"
LINE = "E3DCD0"
WHITE = "FFFFFF"
GREEN = "1B5E45"
GREEN_BG = "E8F3ED"
TEAL = "1E5460"
SLATE = "3E4C59"
ROW = "FBFAF7"

SANS = "Noto Sans"
SANS_TH = "Noto Sans Thai"
SERIF_TH = "Noto Serif Thai"

FONT_FILES = {
    (SANS, False): "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
    (SANS, True): "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf",
    (SANS_TH, False): "/usr/share/fonts/truetype/noto/NotoSansThai-Regular.ttf",
    (SANS_TH, True): "/usr/share/fonts/truetype/noto/NotoSansThai-Bold.ttf",
    (SERIF_TH, False): "/usr/share/fonts/truetype/noto/NotoSerifThai-Regular.ttf",
    (SERIF_TH, True): "/usr/share/fonts/truetype/noto/NotoSerifThai-Bold.ttf",
}


def rgb(hex_color: str) -> RGBColor:
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def tw(length) -> int:
    return int(length.twips)


def content_width(section) -> int:
    return tw(section.page_width) - tw(section.left_margin) - tw(section.right_margin)


def _append(parent, tag, **attrs):
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(key), str(value))
    parent.append(el)
    return el


def prep(p, *, before=0, after=6, line=1.15, align="left", keep_next=False, keep_lines=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p_pr = p._p.get_or_add_pPr()
    _append(p_pr, "w:snapToGrid", **{"w:val": "0"})
    _append(p_pr, "w:autoSpaceDE", **{"w:val": "0"})
    _append(p_pr, "w:autoSpaceDN", **{"w:val": "0"})
    if keep_next:
        _append(p_pr, "w:keepNext", **{"w:val": "true"})
    if keep_lines:
        _append(p_pr, "w:keepLines", **{"w:val": "true"})
    return p


def para(container, *, before=0, after=6, line=1.15, align="left", keep_next=False, keep_lines=False):
    if isinstance(container, _Cell):
        existing = container.paragraphs[0]
        if len(container.paragraphs) == 1 and not existing.text and not existing.runs:
            p = existing
        else:
            p = container.add_paragraph()
    else:
        p = container.add_paragraph()
    return prep(
        p,
        before=before,
        after=after,
        line=line,
        align=align,
        keep_next=keep_next,
        keep_lines=keep_lines,
    )


def run(
    p,
    text,
    *,
    size=11,
    bold=False,
    color=INK,
    thai=SANS_TH,
    latin=SANS,
    tracking=None,
    lang="th-TH",
):
    r = p.add_run(text)
    r.bold = bold
    r.italic = False
    r.font.size = Pt(size)
    r.font.color.rgb = rgb(color)
    r_pr = r._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:cs"), thai)
    r_fonts.set(qn("w:eastAsia"), thai)
    for tag in ("w:sz", "w:szCs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            r_pr.append(node)
        node.set(qn("w:val"), str(int(round(size * 2))))
    lang_el = r_pr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        r_pr.append(lang_el)
    lang_el.set(qn("w:val"), lang)
    lang_el.set(qn("w:eastAsia"), "th-TH")
    lang_el.set(qn("w:bidi"), "th-TH" if lang == "th-TH" else "en-US")
    if tracking is not None:
        _append(r_pr, "w:spacing", **{"w:val": str(tracking)})
    return r


_SCRIPT = re.compile(r"[\u0e00-\u0e7f]+|[^\u0e00-\u0e7f]+")


def write_text(p, text, **kwargs):
    """Keep Thai and Latin in separate runs so English words are not split."""
    if not text:
        return
    for part in _SCRIPT.findall(text):
        lang = "th-TH" if "\u0e00" <= part[0] <= "\u0e7f" else "en-US"
        run(p, part, lang=lang, **kwargs)


def add_text(container, text, **kwargs):
    before = kwargs.pop("before", 0)
    after = kwargs.pop("after", 8)
    line = kwargs.pop("line", 1.15)
    align = kwargs.pop("align", "left")
    keep_next = kwargs.pop("keep_next", False)
    keep_lines = kwargs.pop("keep_lines", False)
    p = para(
        container,
        before=before,
        after=after,
        line=line,
        align=align,
        keep_next=keep_next,
        keep_lines=keep_lines,
    )
    write_text(p, text, **kwargs)
    return p


def border_paragraph(p, *, edge="bottom", color=BRASS, sz="12", space="1"):
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    node = OxmlElement(f"w:{edge}")
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), sz)
    node.set(qn("w:space"), space)
    node.set(qn("w:color"), color)
    p_bdr.append(node)


def shade_cell(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    _append(tc_pr, "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": color})


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:tcMar")):
        tc_pr.remove(old)
    mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def cell_borders(cell, **edges):
    """edges: top=(val, sz, color) or None to leave unset. Use 'nil' to clear."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:tcBorders")):
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = edges.get(edge, ("nil", "0", "FFFFFF"))
        node = OxmlElement(f"w:{edge}")
        val, sz, color = spec
        node.set(qn("w:val"), val)
        node.set(qn("w:sz"), str(sz))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        borders.append(node)
    tc_pr.append(borders)


def set_row_height(row, height, rule="atLeast"):
    tr_pr = row._tr.get_or_add_trPr()
    for old in tr_pr.findall(qn("w:trHeight")):
        tr_pr.remove(old)
    _append(tr_pr, "w:trHeight", **{"w:val": str(height), "w:hRule": rule})


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    _append(tr_pr, "w:cantSplit")


def header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    _append(tr_pr, "w:tblHeader")


def set_table_geometry(table, widths, indent=0):
    table.autofit = False
    table.allow_autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    total = sum(widths)

    def upsert(tag, **attrs):
        node = tbl_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tbl_pr.append(node)
        for key, value in attrs.items():
            node.set(qn(key), str(value))
        return node

    upsert("w:tblW", **{"w:w": total, "w:type": "dxa"})
    upsert("w:tblInd", **{"w:w": indent, "w:type": "dxa"})
    upsert("w:tblLayout", **{"w:type": "fixed"})
    look = tbl_pr.find(qn("w:tblLook"))
    if look is None:
        look = OxmlElement("w:tblLook")
        tbl_pr.append(look)
    look.set(qn("w:val"), "0000")
    look.set(qn("w:firstRow"), "0")
    look.set(qn("w:lastRow"), "0")
    look.set(qn("w:firstColumn"), "0")
    look.set(qn("w:lastColumn"), "0")
    look.set(qn("w:noHBand"), "1")
    look.set(qn("w:noVBand"), "1")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is not None:
        tbl_pr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        node.set(qn("w:sz"), "0")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")
        borders.append(node)
    tbl_pr.append(borders)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        _append(grid, "w:gridCol", **{"w:w": width})
    tbl_pr.addnext(grid)

    for row in table.rows:
        cant_split(row)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def make_table(container, rows, cols, widths):
    table = container.add_table(rows=rows, cols=cols)
    table.style = "Normal Table"
    set_table_geometry(table, widths)
    return table


def v_align(cell, where="center"):
    cell.vertical_alignment = {
        "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
        "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
        "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
    }[where]


def fill_cell(cell, lines, *, fill=None, align="left", valign="center", margins=(70, 70, 110, 110), border=None):
    """lines: list of dicts passed to run(), plus optional before/after/line/break."""
    if fill:
        shade_cell(cell, fill)
    if border:
        cell_borders(cell, **border)
    else:
        cell_borders(cell)
    cell_margins(cell, *margins)
    v_align(cell, valign)
    first = True
    for item in lines:
        item = dict(item)
        before = item.pop("before", 0)
        after = item.pop("after", 0)
        line_sp = item.pop("line", 1.02)
        p = para(cell, before=before, after=after, line=line_sp, align=align, keep_lines=True)
        if item.get("text"):
            write_text(p, item.pop("text"), **item)
        first = False
    return cell


def spacer(container, pt):
    return add_text(container, "", size=2, before=0, after=0, line=1.0).paragraph_format.space_after.__class__ and _spacer(container, pt)


def _spacer(container, pt):
    p = para(container, before=0, after=0, line=1.0)
    p.paragraph_format.space_before = Pt(pt)
    run(p, "", size=2)
    return p


def section_heading(doc, number, title, width):
    p = para(doc, before=0, after=2, line=1.05, keep_next=True, keep_lines=True)
    run(p, number, size=15, bold=True, color=BRASS_TEXT, latin=SANS, thai=SANS_TH)
    run(p, "    ", size=15)
    run(p, title, size=16, bold=True, color=NAVY, thai=SERIF_TH, latin=SANS)
    border_paragraph(p, edge="bottom", color=BRASS, sz="12", space="2")
    p_pr = p._p.get_or_add_pPr()
    _append(p_pr, "w:outlineLvl", **{"w:val": "0"})
    _spacer(doc, 8)
    return p


def subhead(doc, title, *, before=12):
    p = para(doc, before=before, after=4, line=1.05, keep_next=True)
    run(p, title, size=13, bold=True, color=NAVY, thai=SANS_TH, latin=SANS)
    p_pr = p._p.get_or_add_pPr()
    _append(p_pr, "w:outlineLvl", **{"w:val": "1"})
    return p


def body(doc, text, *, after=8, before=0, size=11):
    return add_text(doc, text, size=size, color=INK, before=before, after=after, line=1.18)


def callout(doc, width, text, *, label=None, fill=CREAM, accent=BRASS):
    table = make_table(doc, 1, 1, [width])
    border = {
        "top": ("nil", "0", WHITE),
        "left": ("single", "16", accent),
        "bottom": ("nil", "0", WHITE),
        "right": ("nil", "0", WHITE),
    }
    lines = []
    if label:
        lines.append({"text": label, "size": 9, "bold": True, "color": BRASS_TEXT, "after": 1})
    lines.append({"text": text, "size": 10.5, "color": INK, "line": 1.15})
    fill_cell(table.cell(0, 0), lines, fill=fill, align="left", valign="center", margins=(80, 90, 140, 140), border=border)
    set_row_height(table.rows[0], 360, "atLeast")
    return table


def hairline_rows(table, color=LINE, sz="6", skip_last=False):
    last = len(table.rows) - 1
    for r_idx, row in enumerate(table.rows):
        if skip_last and r_idx == last:
            continue
        for cell in row.cells:
            existing = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
            # reapply bottom while preserving other edges already set
            borders = existing
            if borders is None:
                cell_borders(cell, bottom=("single", sz, color))
            else:
                bottom = borders.find(qn("w:bottom"))
                if bottom is None:
                    bottom = OxmlElement("w:bottom")
                    borders.append(bottom)
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), sz)
                bottom.set(qn("w:space"), "0")
                bottom.set(qn("w:color"), color)


def field_run(p, instruction, *, size=9, bold=False, color=MUTED):
    r = p.add_run()
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = rgb(color)
    r_pr = r._element.get_or_add_rPr()
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), SANS if attr.endswith("ascii") or attr.endswith("hAnsi") else SANS_TH)
    # set explicitly
    r_fonts.set(qn("w:ascii"), SANS)
    r_fonts.set(qn("w:hAnsi"), SANS)
    r_fonts.set(qn("w:cs"), SANS_TH)
    r_fonts.set(qn("w:eastAsia"), SANS_TH)
    r_pr.insert(0, r_fonts)
    for tag in ("w:sz", "w:szCs"):
        node = OxmlElement(tag)
        node.set(qn("w:val"), str(int(size * 2)))
        r_pr.append(node)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1" if instruction.strip() == "PAGE" else "8"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(begin)
    r._r.append(instr)
    r._r.append(separate)
    r._r.append(text)
    r._r.append(end)
    return r


def configure_page(section, *, top, bottom, left, right, header_dist, footer_dist):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    section.header_distance = Cm(header_dist)
    section.footer_distance = Cm(footer_dist)
    section.gutter = Cm(0)


def build_header(section, width):
    header = section.header
    header.is_linked_to_previous = False
    table = header.add_table(1, 2, width=Twips(width))
    table.style = "Normal Table"
    set_table_geometry(table, [int(width * 0.62), width - int(width * 0.62)])
    left, right = table.cell(0, 0), table.cell(0, 1)
    bottom = ("single", "10", BRASS)
    nil = ("nil", "0", WHITE)
    fill_cell(
        left,
        [{"text": "บริษัท ยุธาภัคร์ จำกัด", "size": 9, "bold": True, "color": NAVY}],
        margins=(40, 60, 0, 60),
        border={"top": nil, "left": nil, "right": nil, "bottom": bottom},
        valign="bottom",
    )
    fill_cell(
        right,
        [{"text": "ร่างสำหรับพิจารณาอนุมัติ", "size": 9, "bold": True, "color": BRASS_TEXT}],
        align="right",
        margins=(40, 60, 60, 0),
        border={"top": nil, "left": nil, "right": nil, "bottom": bottom},
        valign="bottom",
    )
    lift_table(header, table)
    p = header.paragraphs[0]
    prep(p, before=0, after=0, line=1.0)
    run(p, "", size=1)


def build_footer(section, width):
    footer = section.footer
    footer.is_linked_to_previous = False
    table = footer.add_table(1, 2, width=Twips(width))
    table.style = "Normal Table"
    set_table_geometry(table, [int(width * 0.72), width - int(width * 0.72)])
    left, right = table.cell(0, 0), table.cell(0, 1)
    top = ("single", "8", BRASS)
    nil = ("nil", "0", WHITE)
    fill_cell(
        left,
        [{"text": "โครงการปรับปรุงระบบการสำรองข้อมูล", "size": 8.5, "color": MUTED}],
        margins=(70, 20, 0, 60),
        border={"top": top, "left": nil, "right": nil, "bottom": nil},
        valign="center",
    )
    # page number cell is built manually so the field sits on one line
    shade = None
    cell_borders(right, top=top, left=nil, right=nil, bottom=nil)
    cell_margins(right, 70, 20, 60, 0)
    v_align(right, "center")
    p = para(right, before=0, after=0, line=1.0, align="right")
    run(p, "หน้า ", size=8.5, color=MUTED)
    field_run(p, "PAGE", size=8.5, color=NAVY, bold=True)
    run(p, "  /  ", size=8.5, color=MUTED)
    field_run(p, "NUMPAGES", size=8.5, color=MUTED)
    lift_table(footer, table)
    p0 = footer.paragraphs[0]
    prep(p0, before=0, after=0, line=1.0)
    run(p0, "", size=1)


def build_cover_footer(section, width):
    footer = section.footer
    footer.is_linked_to_previous = False
    table = footer.add_table(1, 1, width=Twips(width))
    table.style = "Normal Table"
    set_table_geometry(table, [width])
    top = ("single", "8", BRASS)
    nil = ("nil", "0", WHITE)
    fill_cell(
        table.cell(0, 0),
        [{"text": "เอกสารภายใน  ·  ใช้เพื่อพิจารณาอนุมัติ  ·  26 ตุลาคม 2569", "size": 8.5, "color": MUTED}],
        align="center",
        margins=(70, 10, 0, 0),
        border={"top": top, "left": nil, "right": nil, "bottom": nil},
    )
    lift_table(footer, table)
    p0 = footer.paragraphs[0]
    prep(p0, before=0, after=0, line=1.0)
    run(p0, "", size=1)
    # cover has no running header; collapse the default paragraph
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    prep(hp, before=0, after=0, line=1.0)
    run(hp, "", size=1)


def zero_section_break(section):
    sect_pr = section._sectPr
    parent = sect_pr.getparent()
    if parent is not None and parent.tag == qn("w:p"):
        p = parent
        p_pr = p.find(qn("w:pPr"))
        if p_pr is None:
            p_pr = OxmlElement("w:pPr")
            p.insert(0, p_pr)
        spacing = p_pr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            p_pr.append(spacing)
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "20")
        spacing.set(qn("w:lineRule"), "exact")


def metric_row(container, width, items, *, number_size=18, fill=NAVY, number_color=WHITE, label_color=BRASS_ON_NAVY, gap=140):
    n = len(items)
    gaps = n - 1
    usable = width - gap * gaps
    card = usable // n
    widths = []
    for i in range(n):
        widths.append(card)
        if i != n - 1:
            widths.append(gap)
    widths[(n - 1) * 2] += usable - card * n
    table = make_table(container, 1, len(widths), widths)
    nil = ("nil", "0", WHITE)
    item_i = 0
    for idx in range(len(widths)):
        cell = table.cell(0, idx)
        if idx % 2 == 1:
            fill_cell(cell, [{"text": "", "size": 2}], fill=WHITE, margins=(0, 0, 0, 0), border={"top": nil, "left": nil, "right": nil, "bottom": nil})
            continue
        big, label = items[item_i]
        item_i += 1
        fill_cell(
            cell,
            [
                {"text": big, "size": number_size, "bold": True, "color": number_color, "after": 1},
                {"text": label, "size": 8.5, "color": label_color, "line": 1.08},
            ],
            fill=fill,
            align="center",
            valign="center",
            margins=(120, 120, 80, 80),
        )
    set_row_height(table.rows[0], 820, "atLeast")
    return table


def cream_metrics(container, width, items):
    return metric_row(
        container,
        width,
        items,
        number_size=15,
        fill=CREAM,
        number_color=NAVY,
        label_color=MUTED,
        gap=120,
    )


def build_cover(doc, width):
    nil = ("nil", "0", WHITE)
    brass_bottom = ("single", "16", BRASS)
    bar_left = int(width * 0.68)
    bar = make_table(doc, 1, 2, [bar_left, width - bar_left])
    fill_cell(
        bar.cell(0, 0),
        [
            {"text": "บริษัท ยุธาภัคร์ จำกัด", "size": 16, "bold": True, "color": WHITE, "thai": SERIF_TH, "after": 1},
            {"text": "ฝ่ายพัฒนาระบบสารสนเทศ", "size": 10, "color": BRASS_ON_NAVY},
        ],
        fill=NAVY,
        align="left",
        valign="center",
        margins=(120, 130, 180, 80),
        border={"top": nil, "left": nil, "right": nil, "bottom": brass_bottom},
    )
    fill_cell(
        bar.cell(0, 1),
        [
            {"text": "26 ตุลาคม 2569", "size": 11, "bold": True, "color": WHITE, "after": 1},
            {"text": "เอกสารภายใน", "size": 10, "color": BRASS_ON_NAVY},
        ],
        fill=NAVY,
        align="right",
        valign="center",
        margins=(120, 130, 80, 180),
        border={"top": nil, "left": nil, "right": nil, "bottom": brass_bottom},
    )
    set_row_height(bar.rows[0], 980, "atLeast")

    _spacer(doc, 22)
    add_text(
        doc,
        "ENTERPRISE BACKUP SOLUTION",
        size=10.5,
        bold=True,
        color=BRASS_TEXT,
        align="center",
        before=0,
        after=2,
        line=1.0,
        tracking=36,
        latin=SANS,
        thai=SANS_TH,
    )
    add_text(
        doc,
        "โครงการปรับปรุงระบบการสำรองข้อมูล",
        size=26,
        bold=True,
        color=NAVY,
        align="center",
        before=2,
        after=2,
        line=1.05,
        thai=SERIF_TH,
        latin=SANS,
    )
    add_text(
        doc,
        "จากเครื่องแม่ข่าย (PC Server) สู่ระบบจัดเก็บข้อมูลเครือข่าย (NAS)",
        size=12,
        color=SLATE,
        align="center",
        before=0,
        after=2,
        line=1.05,
    )
    p = para(doc, before=2, after=2, line=1.0, align="center")
    run(p, "อัตโนมัติ", size=11, bold=True, color=NAVY)
    run(p, "    ·    ", size=11, color=BRASS)
    run(p, "นอกสถานที่", size=11, bold=True, color=NAVY)
    run(p, "    ·    ", size=11, color=BRASS)
    run(p, "ทดสอบการกู้คืน", size=11, bold=True, color=NAVY)

    _spacer(doc, 8)
    rule_w = 1500
    side = (width - rule_w) // 2
    rule = make_table(doc, 1, 3, [side, rule_w, width - side - rule_w])
    for i, color in ((0, WHITE), (1, BRASS), (2, WHITE)):
        fill_cell(
            rule.cell(0, i),
            [{"text": "", "size": 1}],
            fill=color,
            margins=(0, 0, 0, 0),
        )
    set_row_height(rule.rows[0], 90, "exact")

    _spacer(doc, 10)
    add_text(
        doc,
        "ข้อเสนอจัดหาระบบสำรองข้อมูลอัตโนมัติตามหลัก 3-2-1 เพื่อไม่ให้ข้อมูลสำคัญของบริษัทฯ สูญหายอย่างถาวร และให้กู้คืนระบบได้ภายในเวลาที่กำหนด",
        size=11.5,
        color=INK,
        align="center",
        before=0,
        after=2,
        line=1.2,
    )

    _spacer(doc, 12)
    metric_row(
        doc,
        width,
        [
            ("127,651", "บาท  ·  งบประมาณปีแรก รวมภาษี 7%"),
            ("5,300", "บาทต่อปี  ·  ตั้งแต่ปีที่ 2"),
            ("11 สัปดาห์", "ระยะดำเนินงานหลังได้รับอนุมัติ"),
        ],
        number_size=20,
    )

    _spacer(doc, 14)
    add_text(
        doc,
        "มีข้อมูลอย่างน้อย 3 ชุด  บนอุปกรณ์คนละชนิด  และเก็บไว้นอกสถานที่ 1 ชุด",
        size=10.5,
        color=SLATE,
        align="center",
        before=0,
        after=2,
        line=1.05,
    )

    _spacer(doc, 16)
    rows = [
        ("วันที่จัดทำ", "26 ตุลาคม 2569"),
        ("จัดทำโดย", "ฝ่ายพัฒนาระบบสารสนเทศ"),
        ("เสนอต่อ", "ผู้บริหาร / กรรมการผู้จัดการ"),
        ("สถานะ", "ร่างสำหรับพิจารณาอนุมัติ"),
    ]
    label_w = 2800
    meta = make_table(doc, len(rows), 2, [label_w, width - label_w])
    for i, (label, value) in enumerate(rows):
        value_color = BRASS_TEXT if label == "สถานะ" else NAVY
        fill_cell(
            meta.cell(i, 0),
            [{"text": label, "size": 10.5, "color": MUTED}],
            fill=WHITE,
            margins=(70, 70, 40, 80),
            valign="center",
        )
        fill_cell(
            meta.cell(i, 1),
            [{"text": value, "size": 11.5, "bold": True, "color": value_color}],
            fill=WHITE,
            margins=(70, 70, 80, 40),
            valign="center",
        )
    hairline_rows(meta, color=LINE, sz="6", skip_last=False)
    _spacer(doc, 8)
    add_text(
        doc,
        "วงเงินปีแรกเป็นประมาณการเพื่อขออนุมัติ  ฝ่ายจัดซื้อจะเปรียบเทียบราคาจากผู้จำหน่ายอย่างน้อย 3 รายก่อนสั่งซื้อ",
        size=9,
        color=MUTED,
        align="left",
        before=2,
        after=0,
        line=1.1,
    )


def comparison_table(doc, width):
    headers = ["ประเด็น", "ปัจจุบัน", "เมื่อโครงการเสร็จ"]
    data = [
        (
            "การสำรอง",
            "คัดลอกด้วยมือเป็นครั้งคราว ไม่มีตารางเวลา และไม่มีการตรวจสอบผล",
            "สำรองส่วนเพิ่มทุกวันเวลา 22:00 น. และสำรองทั้งชุดทุกวันอาทิตย์",
        ),
        (
            "นอกสถานที่",
            "ไม่มีสำเนานอกสำนักงาน หากเกิดภัยพิบัติอาจสูญหายทั้งหมด",
            "ส่งขึ้น Cloud Storage ทุกสิ้นเดือน และมีฮาร์ดดิสก์ภายนอกไว้นอกสถานที่",
        ),
        (
            "การตรวจสอบ",
            "ไม่ทราบเมื่อการสำรองล้มเหลว และไม่เคยทดสอบการกู้คืน",
            "แจ้งเตือนเมื่อสำรองไม่สำเร็จ และทดสอบกู้คืนจริงก่อนส่งมอบ",
        ),
        (
            "ความปลอดภัย",
            "สำเนาอยู่ใกล้ข้อมูลต้นทาง เสี่ยงถูกเข้ารหัสพร้อมกันหากถูกโจมตี",
            "แยก VLAN เข้ารหัสข้อมูล จำกัดสิทธิ์ ใช้การยืนยันตัวตนสองขั้นตอน และปิดการเข้าจากภายนอก",
        ),
        (
            "การกู้คืน",
            "ยืนยันไม่ได้ว่าข้อมูลสำรองนำกลับมาใช้ได้",
            "ข้อมูลสูญหายได้ไม่เกิน 24 ชั่วโมง กู้คืนภายใน 4 ชั่วโมงหากอุปกรณ์หลักเสีย และ 48 ชั่วโมงหากเกิดภัยพิบัติ",
        ),
    ]
    widths = [1700, (width - 1700) // 2, width - 1700 - (width - 1700) // 2]
    table = make_table(doc, 1 + len(data), 3, widths)
    nil = ("nil", "0", WHITE)
    for c, text in enumerate(headers):
        fill_cell(
            table.cell(0, c),
            [{"text": text, "size": 10, "bold": True, "color": WHITE}],
            fill=NAVY,
            align="left",
            valign="center",
            margins=(70, 70, 90, 90),
            border={"top": nil, "left": nil, "right": nil, "bottom": nil},
        )
    header_row(table.rows[0])
    for r, row_data in enumerate(data, start=1):
        bg = [CREAM, WHITE, CREAM_DEEP]
        for c, text in enumerate(row_data):
            fill_cell(
                table.cell(r, c),
                [{"text": text, "size": 9.5, "bold": c == 0, "color": NAVY if c == 0 else INK, "line": 1.12}],
                fill=bg[c] if r % 2 else (WHITE if c != 2 else CREAM),
                align="left",
                valign="center",
                margins=(70, 70, 90, 90),
            )
    hairline_rows(table, color=LINE, sz="4")
    # restore header with no hairline clash: header already navy
    for cell in table.rows[0].cells:
        cell_borders(cell)
    return table


def simple_table(doc, width, headers, rows, col_widths, *, aligns=None, emphasize_last=False, font=10):
    aligns = aligns or ["left"] * len(headers)
    table = make_table(doc, 1 + len(rows), len(headers), col_widths)
    nil = ("nil", "0", WHITE)
    for c, text in enumerate(headers):
        fill_cell(
            table.cell(0, c),
            [{"text": text, "size": font, "bold": True, "color": WHITE}],
            fill=NAVY,
            align=aligns[c],
            margins=(64, 64, 90, 90),
        )
    header_row(table.rows[0])
    for r, row_data in enumerate(rows, start=1):
        last = emphasize_last and r == len(rows)
        for c, text in enumerate(row_data):
            fill_cell(
                table.cell(r, c),
                [{
                    "text": text,
                    "size": font + (0.5 if last else 0),
                    "bold": last or c == 0 and False,
                    "color": WHITE if last else INK,
                    "line": 1.08,
                }],
                fill=NAVY if last else (ROW if r % 2 == 0 else WHITE),
                align=aligns[c],
                valign="center",
                margins=(60, 60, 90, 90),
            )
        if not last:
            pass
    if not emphasize_last:
        hairline_rows(table, color=LINE, sz="4")
        for cell in table.rows[0].cells:
            cell_borders(cell)
    else:
        for r_idx, row in enumerate(table.rows):
            if r_idx == 0 or r_idx == len(table.rows) - 1:
                for cell in row.cells:
                    cell_borders(cell)
            else:
                for cell in row.cells:
                    cell_borders(cell, bottom=("single", "4", LINE))
    return table


def build_summary(doc, width):
    section_heading(doc, "01", "บทสรุปผู้บริหาร", width)
    banner = make_table(doc, 1, 1, [width])
    nil = ("nil", "0", WHITE)
    fill_cell(
        banner.cell(0, 0),
        [
            {"text": "ข้อเสนอเพื่อพิจารณาอนุมัติ", "size": 9, "bold": True, "color": BRASS_ON_NAVY, "after": 1},
            {"text": "อนุมัติงบประมาณปีแรก 127,651 บาท", "size": 16, "bold": True, "color": WHITE, "thai": SERIF_TH, "after": 2},
            {"text": "รวมภาษีมูลค่าเพิ่ม 7% แล้ว    ·    รับทราบค่าใช้จ่ายต่อเนื่องปีละ 5,300 บาท ตั้งแต่ปีที่ 2", "size": 10.5, "color": "F3EDE2"},
        ],
        fill=NAVY,
        align="left",
        valign="center",
        margins=(130, 140, 180, 180),
        border={"top": ("single", "12", BRASS), "left": nil, "right": nil, "bottom": nil},
    )
    _spacer(doc, 10)
    body(
        doc,
        "ปัจจุบันบริษัทฯ ยังไม่มีระบบสำรองข้อมูลอัตโนมัติ การสำรองเป็นการคัดลอกด้วยมือไปยังฮาร์ดดิสก์ภายนอกเป็นครั้งคราว ไม่สม่ำเสมอ ไม่มีการตรวจสอบผล และไม่มีสำเนานอกสถานที่ หากอุปกรณ์เสียหาย เกิดภัยพิบัติ หรือถูกโจมตีทางไซเบอร์ ข้อมูลสำคัญอาจสูญหายอย่างถาวร",
    )
    body(
        doc,
        "โครงการนี้จะติดตั้งอุปกรณ์จัดเก็บข้อมูลเครือข่าย (NAS) พร้อมซอฟต์แวร์สำรองระดับองค์กร ออกแบบตามหลัก 3-2-1 คือมีข้อมูลอย่างน้อย 3 ชุด อยู่บนอุปกรณ์คนละชนิด และเก็บไว้นอกสถานที่ 1 ชุด",
        after=8,
    )
    comparison_table(doc, width)
    _spacer(doc, 10)
    cream_metrics(
        doc,
        width,
        [
            ("127,651 บาท", "งบปีแรก รวมภาษี 7%"),
            ("5,300 บาท", "ต่อปี ตั้งแต่ปีที่ 2"),
            ("11 สัปดาห์", "หลังได้รับอนุมัติ"),
            ("4–48 ชม.", "เป้าหมายเวลากู้คืน"),
        ],
    )
    _spacer(doc, 10)
    callout(
        doc,
        width,
        "การลงทุนนี้เป็นการป้องกันความเสี่ยง เมื่อเทียบกับความเสียหายหากข้อมูลสำคัญสูญหายและกู้คืนไม่ได้ ค่ากู้ข้อมูลจากภายนอกเริ่มต้นที่ 30,000 บาท และอาจสูงกว่านั้นมาก นอกจากนี้ยังมีความเสียหายจากการหยุดชะงักของธุรกิจ และความเสี่ยงด้านการปฏิบัติตามกฎหมายคุ้มครองข้อมูลส่วนบุคคล",
    )


def build_problem(doc, width):
    section_heading(doc, "02", "ความเป็นมาและปัญหาที่พบ", width)
    body(
        doc,
        "ข้อมูลของบริษัทฯ จัดเก็บอยู่บน PC Server เพียงเครื่องเดียว การสำรองทำด้วยมือไปยังฮาร์ดดิสก์ภายนอก ไม่มีตารางเวลาที่ชัดเจน และไม่มีการตรวจว่าการสำรองสำเร็จหรือไม่",
    )
    body(
        doc,
        "ปริมาณข้อมูลปัจจุบันประมาณ 880 GB ถึง 1 TB และคาดว่าจะเพิ่มขึ้นร้อยละ 20–30 ต่อปี ประกอบด้วยข้อมูลบัญชี การเงิน ฐานข้อมูลลูกค้า เอกสารโครงการ และข้อมูลส่วนบุคคล",
        after=8,
    )
    cream_metrics(
        doc,
        width,
        [
            ("880 GB–1 TB", "ปริมาณข้อมูลปัจจุบัน"),
            ("20–30%", "อัตราเติบโตที่คาดไว้ต่อปี"),
            ("PDPA", "มีข้อมูลส่วนบุคคลในระบบ"),
        ],
    )
    _spacer(doc, 12)
    subhead(doc, "ปัญหาที่พบ")
    problems = [
        ("ไม่มีระบบอัตโนมัติ", "การสำรองอาศัยคนทำด้วยมือ จึงอาจลืมหรือข้ามขั้นตอน และไม่มีรอบเวลาที่ถือปฏิบัติได้"),
        ("ไม่มีสำเนานอกสถานที่", "หากเกิดภัยพิบัติที่สำนักงาน ข้อมูลต้นทางและฮาร์ดดิสก์ที่เก็บไว้ในที่เดียวกันอาจสูญหายพร้อมกัน"),
        ("ไม่เคยทดสอบการกู้คืน", "ยังไม่มีหลักฐานว่าไฟล์สำรองเปิดใช้งานได้จริง เมื่อเกิดเหตุจึงยืนยันความพร้อมไม่ได้"),
        ("ไม่มีการแจ้งเตือน", "หากการสำรองล้มเหลว จะไม่มีใครทราบ จนกว่าจะต้องใช้ข้อมูลนั้น"),
        ("เสี่ยงต่อการถูกเข้ารหัสเรียกค่าไถ่", "ข้อมูลบนเครื่องที่ใช้งานอาจถูก Ransomware เข้ารหัส หากไม่มีสำเนาที่แยกออกมาและป้องกันไว้ อาจไม่มีชุดข้อมูลที่กู้คืนได้"),
    ]
    label_w = 900
    table = make_table(doc, len(problems), 2, [label_w, width - label_w])
    for i, (title, detail) in enumerate(problems):
        fill_cell(
            table.cell(i, 0),
            [{"text": f"{i + 1:02d}", "size": 12, "bold": True, "color": BRASS_TEXT}],
            fill=CREAM,
            align="center",
            valign="center",
            margins=(80, 80, 40, 40),
        )
        p_lines = [
            {"text": title, "size": 11, "bold": True, "color": NAVY, "after": 1},
            {"text": detail, "size": 10.5, "color": INK, "line": 1.12},
        ]
        fill_cell(
            table.cell(i, 1),
            p_lines,
            fill=WHITE if i % 2 == 0 else ROW,
            align="left",
            valign="center",
            margins=(70, 70, 120, 100),
        )
    _spacer(doc, 10)
    callout(
        doc,
        width,
        "หากยังไม่ดำเนินการ ข้อมูลสำคัญอาจสูญหายอย่างถาวร การทำงานหยุดชะงัก ความน่าเชื่อถือของบริษัทฯ เสียหาย และอาจไม่สอดคล้องกับหน้าที่ตามกฎหมายคุ้มครองข้อมูลส่วนบุคคล",
        label="ผลกระทบหากไม่ดำเนินการ",
        fill="F8F1F1",
        accent="8A4B4B",
    )


def build_scope(doc, width):
    section_heading(doc, "03", "วัตถุประสงค์และขอบเขต", width)
    objectives = [
        "มีระบบสำรองข้อมูลอัตโนมัติที่ทำซ้ำได้และตรวจสอบผลได้",
        "ลดความเสี่ยงจากการสูญเสียข้อมูล ไม่ว่าสาเหตุจะเป็นอุปกรณ์ขัดข้อง ภัยพิบัติ หรือการโจมตี",
        "จัดเก็บสำเนาตามหลัก 3-2-1 รวมถึงสำเนานอกสถานที่",
        "กู้คืนระบบได้รวดเร็ว และลดเวลาที่งานต้องหยุดลง",
        "ยกระดับการคุ้มครองข้อมูลให้สอดคล้องกับกฎหมายและแนวปฏิบัติที่ใช้กันทั่วไป",
    ]
    for i, text in enumerate(objectives, start=1):
        p = para(doc, before=1, after=3, line=1.12, keep_next=True)
        p.paragraph_format.left_indent = Cm(0.85)
        p.paragraph_format.first_line_indent = Cm(-0.85)
        run(p, f"{i:02d}", size=11, bold=True, color=BRASS_TEXT)
        write_text(p, "    " + text, size=11, color=INK)

    _spacer(doc, 8)
    subhead(doc, "ขอบเขตงาน")
    gap = 160
    col = (width - gap) // 2
    widths = [col, gap, width - col - gap]
    included = [
        "จัดหาและติดตั้ง NAS, ฮาร์ดดิสก์, อุปกรณ์เครือข่าย และเครื่องสำรองไฟ",
        "ติดตั้งซอฟต์แวร์สำรองข้อมูล และแยก VLAN สำหรับงานสำรอง",
        "กำหนดนโยบาย ตารางเวลาการสำรอง และระบบแจ้งเตือน",
        "สำรองนอกสถานที่ผ่าน Cloud Storage และฮาร์ดดิสก์ภายนอก",
        "ทดสอบระบบ จัดทำคู่มือ และอบรมผู้ดูแล",
    ]
    excluded = [
        "การเปลี่ยนแปลงหรืออัปเกรด PC Server หลัก",
        "การจัดตั้งศูนย์สำรองฉุกเฉินเต็มรูปแบบ",
    ]
    table = make_table(doc, 2, 3, widths)
    fill_cell(
        table.cell(0, 0),
        [{"text": "รวมอยู่ในโครงการ", "size": 10.5, "bold": True, "color": WHITE}],
        fill=GREEN,
        margins=(70, 70, 100, 100),
    )
    fill_cell(table.cell(0, 1), [{"text": "", "size": 2}], fill=WHITE, margins=(0, 0, 0, 0))
    fill_cell(
        table.cell(0, 2),
        [{"text": "ไม่รวมอยู่ในโครงการ", "size": 10.5, "bold": True, "color": WHITE}],
        fill=SLATE,
        margins=(70, 70, 100, 100),
    )
    fill_cell(
        table.cell(1, 0),
        [{"text": item, "size": 10, "color": INK, "line": 1.12, "after": 4} for item in included],
        fill=GREEN_BG,
        valign="top",
        margins=(80, 70, 110, 110),
    )
    fill_cell(table.cell(1, 1), [{"text": "", "size": 2}], fill=WHITE, margins=(0, 0, 0, 0))
    fill_cell(
        table.cell(1, 2),
        [{"text": item, "size": 10, "color": INK, "line": 1.12, "after": 4} for item in excluded],
        fill="F4F5F6",
        valign="top",
        margins=(80, 70, 110, 110),
    )
    _spacer(doc, 6)
    add_text(
        doc,
        "งานนอกขอบเขตข้างต้นไม่ทำให้เป้าหมายของโครงการนี้เปลี่ยนไป ระบบที่เสนอเป็นการสำรองข้อมูลของเครื่องแม่ข่ายที่มีอยู่ ไม่ใช่การสร้างศูนย์คอมพิวเตอร์สำรอง",
        size=10,
        color=MUTED,
        before=2,
        after=2,
        line=1.12,
    )


def build_architecture(doc, width):
    section_heading(doc, "04", "ระบบที่เสนอและสถาปัตยกรรม", width)
    body(
        doc,
        "แกนของระบบคืออุปกรณ์ NAS แยกจากเครื่องแม่ข่ายที่ใช้งานจริง สำเนาที่สามอยู่บนบริการคลาวด์และฮาร์ดดิสก์ภายนอก จึงครบหลัก 3-2-1",
        after=8,
    )
    metric_row(
        doc,
        width,
        [
            ("3 ชุด", "ข้อมูลใช้งาน  NAS  และสำเนานอกสถานที่"),
            ("2 ชนิด", "ไม่เก็บสำเนาไว้บนอุปกรณ์เดียวกับต้นทาง"),
            ("1 ชุด", "สำเนานอกสถานที่ ผ่าน Cloud และฮาร์ดดิสก์ภายนอก"),
        ],
        number_size=16,
        fill=NAVY,
        number_color=WHITE,
        label_color=BRASS_ON_NAVY,
    )
    page_break(doc)
    subhead(doc, "อุปกรณ์หลัก", before=2)
    specs = [
        ("รุ่น", "Synology DS425+  ขนาด 4 ช่องใส่ดิสก์"),
        ("ฮาร์ดดิสก์", "เกรดสำหรับ NAS  ความจุ 4 TB  จำนวน 4 ลูก"),
        ("การป้องกันดิสก์เสีย", "RAID 5 หรือ RAID 6  ระบบยังทำงานได้แม้มีดิสก์เสียหาย และสร้างข้อมูลขึ้นใหม่ได้"),
        ("เครือข่าย", "VLAN แยกเฉพาะงานสำรอง  ไม่ปะปนกับเครือข่ายใช้งานปกติ"),
        ("ไฟสำรอง", "ติดตั้งเครื่องสำรองไฟ (UPS) สำหรับ NAS และอุปกรณ์เครือข่ายที่เกี่ยวข้อง"),
    ]
    label_w = 2600
    spec = make_table(doc, len(specs), 2, [label_w, width - label_w])
    for i, (label, value) in enumerate(specs):
        fill_cell(
            spec.cell(i, 0),
            [{"text": label, "size": 10.5, "bold": True, "color": NAVY}],
            fill=CREAM,
            margins=(52, 52, 100, 80),
        )
        fill_cell(
            spec.cell(i, 1),
            [{"text": value, "size": 10.5, "color": INK, "line": 1.08}],
            fill=WHITE if i % 2 == 0 else ROW,
            margins=(52, 52, 100, 100),
        )
    _spacer(doc, 8)
    subhead(doc, "ตารางการสำรอง", before=6)
    schedule = [
        ("ทุกวัน  22:00 น.", "สำรองส่วนเพิ่ม", "สำรองเฉพาะข้อมูลที่เปลี่ยนแปลง ใช้เวลาน้อยและประหยัดพื้นที่"),
        ("ทุกวันอาทิตย์", "สำรองทั้งชุด", "สำรองข้อมูลทั้งหมดหนึ่งครั้งต่อสัปดาห์ เป็นจุดกู้คืนที่สมบูรณ์"),
        ("ทุกสิ้นเดือน", "สำเนานอกสถานที่", "ซิงค์ขึ้น Cloud Storage เป็นชุดที่สาม นอกสำนักงาน"),
    ]
    gap = 140
    col = (width - gap * 2) // 3
    widths = [col, gap, col, gap, width - col * 2 - gap * 2]
    sched = make_table(doc, 1, 5, widths)
    colors = [NAVY, TEAL, "6E5324"]
    s_i = 0
    for idx in range(5):
        cell = sched.cell(0, idx)
        if idx % 2 == 1:
            fill_cell(cell, [{"text": "", "size": 2}], fill=WHITE, margins=(0, 0, 0, 0))
            continue
        when, kind, detail = schedule[s_i]
        s_i += 1
        fill_cell(
            cell,
            [
                {"text": when, "size": 9, "bold": True, "color": BRASS_ON_NAVY, "after": 1},
                {"text": kind, "size": 13, "bold": True, "color": WHITE, "after": 2},
                {"text": detail, "size": 9, "color": "F3EDE2", "line": 1.12},
            ],
            fill=colors[s_i - 1],
            align="left",
            valign="top",
            margins=(90, 90, 100, 100),
        )
    set_row_height(sched.rows[0], 900, "atLeast")
    _spacer(doc, 4)
    add_text(
        doc,
        "หากงานสำรองไม่สำเร็จ ระบบจะแจ้งเตือนผู้ดูแล  ไม่ปล่อยให้ความล้มเหลวเงียบจนถึงวันที่ต้องกู้ข้อมูล",
        size=10.5,
        color=INK,
        before=2,
        after=4,
        line=1.12,
    )
    _spacer(doc, 2)
    subhead(doc, "ความปลอดภัยและการเก็บรักษา", before=6)
    sec_items = [
        ("เข้ารหัสทั้งขณะส่งและขณะจัดเก็บ", "ข้อมูลระหว่างทางและข้อมูลที่เก็บบน NAS ถูกเข้ารหัส"),
        ("สิทธิ์เท่าที่จำเป็น", "จำกัดผู้ใช้ที่เข้าถึงชุดสำรองได้ และเปิดการยืนยันตัวตนสองขั้นตอน"),
        ("ไม่เปิดทางเข้าจากอินเทอร์เน็ตโดยตรง", "ลดพื้นที่ที่ผู้โจมตีจะเข้าถึงชุดสำรองจากภายนอก"),
        ("สำเนาอ่านอย่างเดียว", "ใช้เป็นแนวกัน Ransomware ไม่ให้เขียนทับประวัติการสำรองทั้งหมด"),
    ]
    half = (width - 140) // 2
    sec = make_table(doc, 2, 3, [half, 140, width - half - 140])
    for i, (title, detail) in enumerate(sec_items):
        r, side = divmod(i, 2)
        c = 0 if side == 0 else 2
        fill_cell(
            sec.cell(r, c),
            [
                {"text": title, "size": 10.5, "bold": True, "color": NAVY, "after": 1},
                {"text": detail, "size": 10, "color": INK, "line": 1.1},
            ],
            fill=CREAM if i % 2 == 0 else WHITE,
            margins=(64, 64, 100, 100),
            border={"top": ("nil", "0", WHITE), "left": ("single", "12", BRASS), "bottom": ("nil", "0", WHITE), "right": ("nil", "0", WHITE)},
        )
    fill_cell(sec.cell(0, 1), [{"text": "", "size": 2}], margins=(0, 0, 0, 0))
    fill_cell(sec.cell(1, 1), [{"text": "", "size": 2}], margins=(0, 0, 0, 0))

    _spacer(doc, 6)
    subhead(doc, "เป้าหมายการกู้คืนและการเก็บประวัติ", before=6)
    cream_metrics(
        doc,
        width,
        [
            ("24 ชั่วโมง", "ข้อมูลสูญหายได้ไม่เกินนี้ (RPO)"),
            ("4 ชั่วโมง", "กู้คืนหากฮาร์ดแวร์หลักเสีย (RTO)"),
            ("48 ชั่วโมง", "กู้คืนหากเกิดภัยพิบัติ (RTO)"),
        ],
    )
    _spacer(doc, 8)
    ret_w = width
    each = ret_w // 3
    retention = make_table(doc, 2, 3, [each, each, ret_w - each * 2])
    labels = [("14 วัน", "สำรองรายวัน"), ("8 สัปดาห์", "สำรองรายสัปดาห์"), ("12 เดือน", "สำรองรายเดือน")]
    for c, (num, label) in enumerate(labels):
        fill_cell(
            retention.cell(0, c),
            [{"text": num, "size": 14, "bold": True, "color": WHITE}],
            fill=NAVY_MID,
            align="center",
            margins=(70, 20, 60, 60),
        )
        fill_cell(
            retention.cell(1, c),
            [{"text": label, "size": 9.5, "color": NAVY}],
            fill=CREAM,
            align="center",
            margins=(20, 70, 60, 60),
        )
    _spacer(doc, 8)
    add_text(
        doc,
        "ก่อนส่งมอบจะทดสอบการกู้คืนกับข้อมูลจริง จัดทำคู่มือปฏิบัติ และอบรมผู้รับผิดชอบมากกว่าหนึ่งคน เพื่อไม่ให้ความรู้ในการกู้ระบบอยู่ที่บุคคลเดียว",
        size=10.5,
        color=INK,
        before=2,
        after=2,
        line=1.15,
    )


def build_budget(doc, width):
    section_heading(doc, "05", "งบประมาณและการวิเคราะห์ค่าใช้จ่าย", width)
    body(
        doc,
        "วงเงินที่ขออนุมัติคืองบปีแรก 127,651 บาท รวมภาษีมูลค่าเพิ่ม 7% แล้ว ตัวเลขนี้เป็นประมาณการเพื่อตั้งวงเงิน ฝ่ายจัดซื้อจะขอใบเสนอราคาจากผู้จำหน่ายอย่างน้อย 3 รายก่อนสั่งซื้อ",
        after=8,
    )
    headers = ["รายการ", "จำนวนเงิน (บาท)"]
    rows = [
        ("ฮาร์ดแวร์  —  NAS, ฮาร์ดดิสก์, UPS, สวิตช์ และอุปกรณ์เครือข่าย", "67,200"),
        ("ซอฟต์แวร์และบริการ Cloud ปีแรก", "5,300"),
        ("อุปกรณ์สำรองนอกสถานที่  —  ฮาร์ดดิสก์ภายนอก", "7,800"),
        ("ติดตั้ง ตั้งค่า อบรม และจัดทำเอกสาร", "39,000"),
        ("รวมก่อนภาษี", "119,300"),
        ("ภาษีมูลค่าเพิ่ม 7%", "8,351"),
        ("รวมทั้งสิ้นปีแรก", "127,651"),
    ]
    col_w = [int(width * 0.74), width - int(width * 0.74)]
    table = make_table(doc, 1 + len(rows), 2, col_w)
    nil = ("nil", "0", WHITE)
    for c, text in enumerate(headers):
        fill_cell(
            table.cell(0, c),
            [{"text": text, "size": 10, "bold": True, "color": WHITE}],
            fill=NAVY,
            align="left" if c == 0 else "right",
            margins=(70, 70, 100, 100),
        )
    header_row(table.rows[0])
    for r, (name, amount) in enumerate(rows, start=1):
        kind = "normal"
        if r == len(rows):
            kind = "total"
        elif r >= len(rows) - 2:
            kind = "sub"
        fill = NAVY if kind == "total" else (CREAM if kind == "sub" else (ROW if r % 2 == 0 else WHITE))
        color = WHITE if kind == "total" else NAVY
        name_size = 11 if kind == "total" else 10.5
        fill_cell(
            table.cell(r, 0),
            [{"text": name, "size": name_size, "bold": kind != "normal", "color": color}],
            fill=fill,
            align="left",
            margins=(64, 64, 100, 80),
        )
        fill_cell(
            table.cell(r, 1),
            [{"text": amount, "size": name_size, "bold": True, "color": color}],
            fill=fill,
            align="right",
            margins=(64, 64, 80, 110),
        )
        if kind == "normal":
            for cell in table.rows[r].cells:
                cell_borders(cell, bottom=("single", "4", LINE))
        else:
            for cell in table.rows[r].cells:
                cell_borders(cell)
    for cell in table.rows[0].cells:
        cell_borders(cell)

    _spacer(doc, 8)
    add_text(doc, "สัดส่วนค่าใช้จ่ายก่อนภาษี 119,300 บาท", size=10, bold=True, color=NAVY, before=2, after=4, keep_next=True)

    parts = [
        (67200, NAVY, "ฮาร์ดแวร์  67,200"),
        (39000, BRASS, "ติดตั้ง อบรม เอกสาร  39,000"),
        (7800, TEAL, "ดิสก์นอกสถานที่  7,800"),
        (5300, "8C7A62", "ซอฟต์แวร์และคลาวด์  5,300"),
    ]
    total_amt = sum(p[0] for p in parts)
    raw = [width * amt / total_amt for _, amt in [(None, amt) for amt, _, _ in parts]]
    # recompute properly
    raw = [width * amt / total_amt for amt, _, _ in parts]
    bar_widths = [int(x) for x in raw]
    bar_widths[-1] += width - sum(bar_widths)
    bar = make_table(doc, 1, 4, bar_widths)
    for i, (_, color, _) in enumerate(parts):
        fill_cell(
            bar.cell(0, i),
            [{"text": "", "size": 2}],
            fill=color,
            margins=(0, 0, 0, 0),
        )
    set_row_height(bar.rows[0], 160, "exact")
    _spacer(doc, 4)
    legend = make_table(doc, 1, 4, [width // 4] * 3 + [width - (width // 4) * 3])
    for i, (_, color, label) in enumerate(parts):
        fill_cell(
            legend.cell(0, i),
            [{"text": label, "size": 8, "bold": True, "color": NAVY}],
            fill=WHITE,
            align="left",
            valign="center",
            margins=(20, 20, 40, 40),
            border={"top": ("single", "16", color), "left": ("nil", "0", WHITE), "right": ("nil", "0", WHITE), "bottom": ("nil", "0", WHITE)},
        )

    _spacer(doc, 10)
    gap = 160
    half = (width - gap) // 2
    pair = make_table(doc, 1, 3, [half, gap, width - half - gap])
    fill_cell(
        pair.cell(0, 0),
        [
            {"text": "ตั้งแต่ปีที่ 2", "size": 9, "bold": True, "color": BRASS_ON_NAVY, "after": 1},
            {"text": "5,300 บาทต่อปี", "size": 16, "bold": True, "color": WHITE, "after": 2},
            {"text": "ค่าบริการ Cloud และสิทธิ์การใช้ซอฟต์แวร์ ไม่มีค่างวดฮาร์ดแวร์ซ้ำ", "size": 9.5, "color": "F3EDE2", "line": 1.1},
        ],
        fill=NAVY,
        margins=(120, 120, 140, 140),
    )
    fill_cell(pair.cell(0, 1), [{"text": "", "size": 2}], margins=(0, 0, 0, 0))
    fill_cell(
        pair.cell(0, 2),
        [
            {"text": "รวม 3 ปี ตามรายการนี้", "size": 9, "bold": True, "color": BRASS_TEXT, "after": 1},
            {"text": "138,251 บาท", "size": 16, "bold": True, "color": NAVY, "after": 2},
            {"text": "ปีแรก 127,651 บวกปีที่ 2 และปีที่ 3 ปีละ 5,300", "size": 9.5, "color": SLATE, "line": 1.1},
        ],
        fill=CREAM,
        margins=(120, 120, 140, 140),
    )

    _spacer(doc, 10)
    subhead(doc, "สิ่งที่วงเงินนี้กันไว้")
    points = [
        ("ค่ากู้ข้อมูลจากภายนอก", "เริ่มต้น 30,000 บาท และอาจสูงกว่ามากหากข้อมูลเสียหายหนักหรือกู้ได้ไม่ครบ"),
        ("การหยุดชะงักของงาน", "งานบัญชี การเงิน ฐานข้อมูลลูกค้า และเอกสารโครงการจะหยุดหากเครื่องหลักใช้การไม่ได้และไม่มีสำเนา"),
        ("หน้าที่ตามกฎหมาย", "ข้อมูลส่วนบุคคลอยู่ในระบบนี้ การไม่มีมาตรการสำรองที่เหมาะสมเพิ่มความเสี่ยงหากข้อมูลสูญหายหรือรั่วไหล"),
    ]
    for title, detail in points:
        p = para(doc, before=2, after=3, line=1.12)
        p.paragraph_format.left_indent = Cm(0.35)
        run(p, title, size=11, bold=True, color=NAVY)
        write_text(p, "  —  " + detail, size=11, color=INK)
    _spacer(doc, 4)
    callout(
        doc,
        width,
        "โครงการนี้ไม่ได้เสนอเพื่อสร้างรายได้ เสนอเพื่อกันความเสียหายที่กู้คืนไม่ได้ เมื่อเทียบกับวงเงินปีแรก 127,651 บาท การมีสำเนาที่ทดสอบแล้วและมีชุดนอกสถานที่ เป็นค่าป้องกันที่ตรวจสอบตัวเลขได้",
    )


def build_plan(doc, width):
    section_heading(doc, "06", "แผนดำเนินงานและการบริหารความเสี่ยง", width)
    body(doc, "ระยะเวลารวมประมาณ 11 สัปดาห์ นับจากวันที่ได้รับอนุมัติและเริ่มจัดซื้อได้", after=8)
    headers = ["ช่วง", "งาน", "สิ่งที่ต้องส่งมอบ"]
    rows = [
        ("สัปดาห์ 1–2", "ขออนุมัติและจัดซื้อจัดจ้าง", "วงเงินอนุมัติ และใบเสนอราคาอย่างน้อย 3 ราย"),
        ("สัปดาห์ 3–5", "ติดตั้งฮาร์ดแวร์และตั้งค่าเครือข่าย", "NAS, ดิสก์, UPS, สวิตช์ และ VLAN พร้อมใช้"),
        ("สัปดาห์ 6–7", "ติดตั้งซอฟต์แวร์และกำหนดนโยบาย", "ตารางสำรอง การแจ้งเตือน และการเข้ารหัสทำงานจริง"),
        ("สัปดาห์ 8–9", "ทดสอบระบบและการกู้คืนข้อมูล", "ผลทดสอบว่านำข้อมูลกลับมาใช้งานได้"),
        ("สัปดาห์ 10–11", "อบรมและส่งมอบงาน", "คู่มือ และผู้ดูแลที่ปฏิบัติได้มากกว่า 1 คน"),
    ]
    widths = [2100, 4300, width - 6400]
    simple_table(doc, width, headers, rows, widths, font=10)
    _spacer(doc, 8)
    subhead(doc, "ความเสี่ยงหลักและทางรับมือ")
    risks = [
        ("ฮาร์ดดิสก์เสียหาย", "ใช้ RAID เพื่อให้ระบบทำงานต่อได้ และสร้างข้อมูลขึ้นใหม่จากดิสก์ที่เหลือ"),
        ("ไฟฟ้าดับ", "ติดตั้งเครื่องสำรองไฟให้ปิดระบบได้อย่างปลอดภัย"),
        ("ถูกโจมตีเรียกค่าไถ่", "แยกเครือข่ายสำรองออกจากเครือข่ายใช้งาน และเก็บภาพสำรองแบบอ่านอย่างเดียว"),
        ("ผู้ดูแลไม่พอหรือลาออก", "ทำคู่มือและอบรมผู้รับผิดชอบมากกว่าหนึ่งคน ขั้นตอนกู้คืนไม่ผูกกับบุคคลเดียว"),
    ]
    risk = make_table(doc, 1 + len(risks), 2, [3200, width - 3200])
    for c, text in enumerate(["ความเสี่ยง", "ทางรับมือ"]):
        fill_cell(
            risk.cell(0, c),
            [{"text": text, "size": 10, "bold": True, "color": WHITE}],
            fill=NAVY,
            margins=(64, 64, 100, 100),
        )
    header_row(risk.rows[0])
    for i, (a, b) in enumerate(risks):
        fill_cell(
            risk.cell(i + 1, 0),
            [{"text": a, "size": 10.5, "bold": True, "color": NAVY}],
            fill=CREAM if i % 2 == 0 else WHITE,
            margins=(64, 64, 100, 80),
        )
        fill_cell(
            risk.cell(i + 1, 1),
            [{"text": b, "size": 10.5, "color": INK, "line": 1.1}],
            fill=ROW if i % 2 else WHITE,
            margins=(64, 64, 80, 100),
        )
    hairline_rows(risk, color=LINE, sz="4")
    for cell in risk.rows[0].cells:
        cell_borders(cell)
    _spacer(doc, 6)
    add_text(
        doc,
        "ให้ทบทวนปริมาณข้อมูลและแผนขยายความจุทุก 6 เดือน อัตราเติบโตที่คาดไว้ร้อยละ 20–30 ต่อปีจะทำให้ความจุที่เผื่อไว้วันนี้ค่อย ๆ ลดลง",
        size=10.5,
        color=INK,
        before=4,
        after=2,
        line=1.15,
    )
    _spacer(doc, 8)
    callout(
        doc,
        width,
        "หากผู้บริหารอนุมัติข้อเสนอในส่วนถัดไป บริษัทฯ จะได้ระบบสำรองอัตโนมัติภายในประมาณ 11 สัปดาห์ ด้วยวงเงินปีแรก 127,651 บาท มีสำเนานอกสถานที่ และมีผลการทดสอบว่านำข้อมูลกลับมาใช้ได้จริง",
        label="หากได้รับอนุมัติ",
    )


def build_approval(doc, width):
    section_heading(doc, "07", "ข้อเสนอเพื่อพิจารณาอนุมัติ", width)
    body(doc, "ฝ่ายพัฒนาระบบสารสนเทศขอให้ผู้บริหารพิจารณา 4 เรื่องดังนี้", after=6)
    items = [
        ("อนุมัติวงเงิน", "อนุมัติงบประมาณ 127,651 บาท สำหรับโครงการปรับปรุงระบบสำรองข้อมูลในปีแรก รวมภาษีมูลค่าเพิ่ม 7% แล้ว"),
        ("เปรียบเทียบราคา", "มอบหมายฝ่ายจัดซื้อขอใบเสนอราคาจากผู้จำหน่ายอย่างน้อย 3 ราย ตามเกณฑ์ที่กำหนด ก่อนผูกพันสั่งซื้อ"),
        ("ผู้รับผิดชอบ", "แต่งตั้งผู้รับผิดชอบโครงการ เพื่อควบคุมให้งานติดตั้ง ทดสอบ และส่งมอบเป็นไปตามแผน 11 สัปดาห์"),
        ("ทบทวนทุก 6 เดือน", "ทบทวนปริมาณข้อมูลและแผนขยายระบบทุก 6 เดือน เพื่อรองรับการเติบโตของข้อมูล"),
    ]
    table = make_table(doc, len(items), 2, [1500, width - 1500])
    for i, (title, detail) in enumerate(items):
        fill_cell(
            table.cell(i, 0),
            [{"text": f"0{i + 1}", "size": 14, "bold": True, "color": WHITE}],
            fill=NAVY,
            align="center",
            margins=(80, 80, 40, 40),
        )
        fill_cell(
            table.cell(i, 1),
            [
                {"text": title, "size": 11, "bold": True, "color": NAVY, "after": 1},
                {"text": detail, "size": 10.5, "color": INK, "line": 1.12},
            ],
            fill=CREAM if i % 2 == 0 else WHITE,
            margins=(70, 70, 120, 100),
        )
    _spacer(doc, 12)
    add_text(
        doc,
        "จึงเรียนมาเพื่อโปรดพิจารณาอนุมัติ",
        size=12,
        bold=True,
        color=NAVY,
        thai=SERIF_TH,
        before=4,
        after=2,
        keep_next=True,
    )

    _spacer(doc, 8)
    gap = 400
    half = (width - gap) // 2
    sig = make_table(doc, 1, 3, [half, gap, width - half - gap])
    blocks = [
        ("ผู้จัดทำ", "ฝ่ายพัฒนาระบบสารสนเทศ", "ผู้เสนอโครงการ"),
        ("ผู้อนุมัติ", "ผู้บริหาร / กรรมการผู้จัดการ", "ผู้มีอำนาจอนุมัติ"),
    ]
    for col, (role, who, note) in zip((0, 2), blocks):
        cell = sig.cell(0, col)
        cell_borders(cell)
        cell_margins(cell, 40, 40, 80, 80)
        v_align(cell, "top")
        shade_cell(cell, WHITE)
        p_role = para(cell, before=0, after=0, line=1.0, keep_next=True, keep_lines=True)
        run(p_role, role, size=10, bold=True, color=BRASS_TEXT)
        p_space = para(cell, before=36, after=0, line=1.0, keep_next=True, keep_lines=True)
        run(p_space, " ", size=11)
        border_paragraph(p_space, edge="bottom", color=NAVY, sz="8", space="1")
        p_who = para(cell, before=4, after=0, line=1.05, keep_next=True)
        run(p_who, who, size=11, bold=True, color=NAVY)
        p_note = para(cell, before=0, after=0, line=1.0, keep_next=True)
        run(p_note, note, size=9, color=MUTED)
        p_date = para(cell, before=8, after=0, line=1.0, keep_lines=True)
        run(p_date, "วันที่ ................................", size=10, color=SLATE)
    fill_cell(sig.cell(0, 1), [{"text": "", "size": 2}], margins=(0, 0, 0, 0))
    _spacer(doc, 12)
    add_text(
        doc,
        "เอกสารฉบับนี้เป็นร่างสำหรับพิจารณาอนุมัติ จัดทำโดยฝ่ายพัฒนาระบบสารสนเทศ เมื่อวันที่ 26 ตุลาคม 2569 รายละเอียดงบประมาณจะยืนยันอีกครั้งจากใบเสนอราคา",
        size=8.5,
        color=MUTED,
        before=4,
        after=0,
        line=1.1,
    )


def apply_doc_defaults(doc):
    styles = doc.styles.element
    existing = styles.find(qn("w:docDefaults"))
    if existing is not None:
        styles.remove(existing)
    defaults = OxmlElement("w:docDefaults")
    r_pr_default = OxmlElement("w:rPrDefault")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), SANS)
    r_fonts.set(qn("w:hAnsi"), SANS)
    r_fonts.set(qn("w:cs"), SANS_TH)
    r_fonts.set(qn("w:eastAsia"), SANS_TH)
    r_pr.append(r_fonts)
    for tag, val in (("w:sz", "22"), ("w:szCs", "22")):
        node = OxmlElement(tag)
        node.set(qn("w:val"), val)
        r_pr.append(node)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "th-TH")
    lang.set(qn("w:eastAsia"), "th-TH")
    lang.set(qn("w:bidi"), "th-TH")
    r_pr.append(lang)
    r_pr_default.append(r_pr)
    defaults.append(r_pr_default)
    p_pr_default = OxmlElement("w:pPrDefault")
    p_pr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "276")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    p_pr_default.append(p_pr)
    defaults.append(p_pr_default)
    styles.insert(0, defaults)

    normal = doc.styles["Normal"]
    normal.font.name = SANS
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    r_pr = normal.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), SANS)
    r_fonts.set(qn("w:hAnsi"), SANS)
    r_fonts.set(qn("w:cs"), SANS_TH)
    r_fonts.set(qn("w:eastAsia"), SANS_TH)


def page_break(doc):
    p = para(doc, before=0, after=0, line=1.0)
    run(p, "", size=2)
    p.runs[0].add_break(WD_BREAK.PAGE)


def build():
    doc = Document()
    apply_doc_defaults(doc)
    core = doc.core_properties
    core.title = "โครงการปรับปรุงระบบการสำรองข้อมูล"
    core.subject = "ข้อเสนอระบบสำรองข้อมูลจาก PC Server สู่ NAS เพื่อพิจารณาอนุมัติ"
    core.category = "ร่างสำหรับพิจารณาอนุมัติ"
    core.keywords = "สำรองข้อมูล; NAS; 3-2-1; PDPA"
    core.language = "th-TH"
    core.creator = "ฝ่ายพัฒนาระบบสารสนเทศ บริษัท ยุธาภัคร์ จำกัด"
    core.author = "ฝ่ายพัฒนาระบบสารสนเทศ บริษัท ยุธาภัคร์ จำกัด"
    core.last_modified_by = "ฝ่ายพัฒนาระบบสารสนเทศ บริษัท ยุธาภัคร์ จำกัด"

    cover = doc.sections[0]
    configure_page(cover, top=1.25, bottom=1.45, left=1.7, right=1.7, header_dist=0.4, footer_dist=0.45)
    width = content_width(cover)
    build_cover_footer(cover, width)
    build_cover(doc, width)

    body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    zero_section_break(cover)
    configure_page(body_sec, top=1.85, bottom=1.55, left=1.7, right=1.7, header_dist=0.55, footer_dist=0.42)
    width = content_width(body_sec)
    build_header(body_sec, width)
    build_footer(body_sec, width)

    build_summary(doc, width)
    page_break(doc)
    build_problem(doc, width)
    page_break(doc)
    build_scope(doc, width)
    build_architecture(doc, width)
    page_break(doc)
    build_budget(doc, width)
    page_break(doc)
    build_plan(doc, width)
    page_break(doc)
    build_approval(doc, width)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    embed_fonts(OUT)
    validate_glyphs(doc)
    return OUT


def _guid_key(guid: str) -> bytes:
    parts = guid.strip("{}").split("-")
    key = bytearray()
    key += bytes.fromhex(parts[0])[::-1]
    key += bytes.fromhex(parts[1])[::-1]
    key += bytes.fromhex(parts[2])[::-1]
    key += bytes.fromhex(parts[3])
    key += bytes.fromhex(parts[4])
    if len(key) != 16:
        raise RuntimeError("bad font GUID key")
    return bytes(key)


def obfuscate(data: bytes, guid: str) -> bytes:
    key = _guid_key(guid)
    buf = bytearray(data)
    for i in range(min(32, len(buf))):
        buf[i] ^= key[i % 16]
    return bytes(buf)


def embed_fonts(path: str) -> None:
    families = {}
    rels = []
    blobs = {}
    rid = 1
    for (family, bold), font_path in FONT_FILES.items():
        guid = "{" + str(uuid.uuid4()).upper() + "}"
        data = obfuscate(open(font_path, "rb").read(), guid)
        name = f"word/fonts/font{rid}.odttf"
        blobs[name] = data
        rel_id = f"rIdFont{rid}"
        rels.append((rel_id, name))
        families.setdefault(family, {})["bold" if bold else "regular"] = (rel_id, guid)
        # sanity: deobfuscate magic
        magic = bytearray(data[:4])
        key = _guid_key(guid)
        for i in range(4):
            magic[i] ^= key[i % 16]
        if bytes(magic) not in (b"\x00\x01\x00\x00", b"OTTO", b"true", b"typ1"):
            raise RuntimeError(f"font obfuscation failed for {family}")
        rid += 1

    font_xml = [
        "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>",
        '<w:fonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">',
    ]
    for family, faces in families.items():
        font_xml.append(f'  <w:font w:name="{escape(family)}">')
        font_xml.append('    <w:charset w:val="00"/>')
        font_xml.append('    <w:family w:val="swiss"/>')
        font_xml.append('    <w:pitch w:val="variable"/>')
        if "regular" in faces:
            rel_id, guid = faces["regular"]
            font_xml.append(f'    <w:embedRegular r:id="{rel_id}" w:fontKey="{guid}"/>')
        if "bold" in faces:
            rel_id, guid = faces["bold"]
            font_xml.append(f'    <w:embedBold r:id="{rel_id}" w:fontKey="{guid}"/>')
        font_xml.append("  </w:font>")
    font_xml.append("</w:fonts>")
    font_table = "\n".join(font_xml).encode("utf-8")

    rels_xml = [
        "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>",
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">',
    ]
    for rel_id, name in rels:
        target = name.split("/", 1)[1]  # fonts/fontN.odttf relative to word/
        rels_xml.append(
            f'  <Relationship Id="{rel_id}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="{target}"/>'
        )
    rels_xml.append("</Relationships>")

    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/fontTable.xml":
                data = font_table
            elif item.filename == "word/settings.xml":
                data = patch_settings(data)
            elif item.filename == "[Content_Types].xml":
                data = patch_content_types(data)
            elif item.filename == "word/_rels/fontTable.xml.rels":
                continue
            zout.writestr(item.filename, data)
        zout.writestr("word/_rels/fontTable.xml.rels", "\n".join(rels_xml).encode("utf-8"))
        for name, blob in blobs.items():
            zout.writestr(name, blob)
    os.replace(tmp, path)


def patch_settings(data: bytes) -> bytes:
    text = data.decode("utf-8")
    text = text.replace("<w:useFELayout/>", "")
    text = text.replace(
        '<w:themeFontLang w:val="en-US" w:eastAsia="ja-JP"/>',
        '<w:themeFontLang w:val="th-TH" w:eastAsia="th-TH" w:bidi="th-TH"/>',
    )
    text = text.replace('<w:zoom w:val="bestFit"/>', '<w:zoom w:percent="110"/>')
    if "<w:embedTrueTypeFonts" not in text:
        text = text.replace(
            "<w:compat>",
            "<w:embedTrueTypeFonts/><w:embedSystemFonts/><w:autoHyphenation w:val=\"false\"/><w:updateFields w:val=\"true\"/><w:compat>",
        )
    return text.encode("utf-8")


def patch_content_types(data: bytes) -> bytes:
    text = data.decode("utf-8")
    if 'Extension="odttf"' not in text:
        insert = '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/>'
        start = text.find("<Types")
        idx = text.find(">", start)
        text = text[: idx + 1] + insert + text[idx + 1 :]
    return text.encode("utf-8")


def lift_table(container, table):
    """Place a header/footer table above the default empty paragraph."""
    tbl = table._tbl
    first_p = container.paragraphs[0]._element
    parent = first_p.getparent()
    parent.remove(tbl)
    first_p.addprevious(tbl)


def validate_glyphs(doc: Document) -> None:
    texts = []

    def walk(paragraphs):
        for p in paragraphs:
            texts.append(p.text)

    walk(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                walk(cell.paragraphs)
    blob = "\n".join(texts)
    needed = {}
    for ch in set(blob):
        if ch.isspace():
            continue
        needed.setdefault("thai" if "\u0e00" <= ch <= "\u0e7f" else "latin", set()).add(ch)
    thai = TTFont(FONT_FILES[(SANS_TH, False)]).getBestCmap()
    latin = TTFont(FONT_FILES[(SANS, False)]).getBestCmap()
    serif = TTFont(FONT_FILES[(SERIF_TH, False)]).getBestCmap()
    missing_th = sorted(ch for ch in needed.get("thai", []) if ord(ch) not in thai or ord(ch) not in serif)
    missing_lat = sorted(ch for ch in needed.get("latin", []) if ord(ch) not in latin)
    if missing_th or missing_lat:
        raise SystemExit(f"Missing glyphs Thai={missing_th} Latin={''.join(missing_lat)}")


if __name__ == "__main__":
    path = build()
    print(path)
